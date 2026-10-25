import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException, status
from pydantic import ValidationError

from app.schemas.dataset_ingestion.dataset_import import DatasetImportPayload

MAX_DATASET_FILE_BYTES = 10 * 1024 * 1024
MAX_DATASET_CASES = 50_000


def _tabular_rows(filename: str, content: bytes) -> list[list[Any]]:
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Supported dataset files are .json, .csv, and .xlsx.",
        )
    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty.",
        )
    if len(content) > MAX_DATASET_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Dataset files must be 10 MB or smaller.",
        )

    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
            rows = []
            for row in csv.reader(io.StringIO(text, newline="")):
                rows.append(row)
                if len(rows) > MAX_DATASET_CASES + 1:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"Dataset files may contain at most {MAX_DATASET_CASES:,} cases.",
                    )
            return rows
        except (UnicodeDecodeError, csv.Error) as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unable to read CSV. Save it as UTF-8 and try again.",
            ) from exc

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = []
            for row in workbook.active.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_DATASET_CASES + 1:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"Dataset files may contain at most {MAX_DATASET_CASES:,} cases.",
                    )
            return rows
        finally:
            workbook.close()
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Excel import support is not installed on the server.",
        ) from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unable to read Excel workbook. Upload a valid .xlsx file.",
        ) from exc


def parse_tabular_dataset(filename: str, content: bytes) -> DatasetImportPayload:
    rows = _tabular_rows(filename, content)
    if not rows:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The file must contain a header row and at least one case.",
        )

    headers = [str(value or "").strip().lower().replace(" ", "_") for value in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or "input" not in headers:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The header row must include an 'input' column.",
        )
    if len(set(headers)) != len(headers):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Column names must be unique after trimming and normalization.",
        )

    cases: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row))
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        record = dict(zip(headers, values))
        raw_input = record.get("input")
        if raw_input is None or not str(raw_input).strip():
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Row {row_number} must contain a non-empty input.",
            )

        metadata: dict[str, Any] = {}
        raw_metadata = record.get("metadata")
        if raw_metadata is not None and str(raw_metadata).strip():
            try:
                metadata_value = (
                    raw_metadata
                    if isinstance(raw_metadata, dict)
                    else json.loads(str(raw_metadata))
                )
            except json.JSONDecodeError as exc:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"Row {row_number} metadata must be a JSON object.",
                ) from exc
            if not isinstance(metadata_value, dict):
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"Row {row_number} metadata must be a JSON object.",
                )
            metadata.update(metadata_value)

        for header, value in record.items():
            if (
                header not in {"input", "expected_output", "metadata"}
                and value is not None
                and value != ""
            ):
                metadata[header] = value.isoformat() if isinstance(value, (date, datetime)) else value

        raw_expected = record.get("expected_output")
        case: dict[str, Any] = {
            "input": str(raw_input),
            "expected_output": None if raw_expected is None else str(raw_expected),
        }
        if metadata:
            case["metadata"] = metadata
        cases.append(case)
        if len(cases) > MAX_DATASET_CASES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"Dataset files may contain at most {MAX_DATASET_CASES:,} cases.",
            )

    try:
        return DatasetImportPayload.model_validate({"cases": cases})
    except ValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Dataset file validation failed.",
                "errors": exc.errors(),
            },
        ) from exc


async def parse_json_dataset(file) -> DatasetImportPayload:
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="A dataset file is required.",
        )

    if not file.filename.lower().endswith(".json"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only JSON dataset files are supported.",
        )

    try:
        content = await file.read()
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unable to read the uploaded file.",
        )

    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded JSON file is empty.",
        )

    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid JSON file: {exc.msg}.",
        )

    try:
        return DatasetImportPayload.model_validate(payload)
    except ValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Dataset JSON validation failed.",
                "errors": exc.errors(),
            },
        )
